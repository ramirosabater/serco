#!/usr/bin/env python3
"""
SGRP 2.0 · Informe de validación de la migración (hito 1).

Compara las cifras de control del Access (esperado-*.json, lo arma Claude con el .mdb)
contra lo que hay en Supabase, y deja el informe en un Excel para entregarle a Serco.

Necesita Python 3 y dos librerías:
    pip install psycopg2-binary openpyxl

Uso (en esta carpeta):
    python validar.py esperado.json "postgresql://postgres.xxxx:CLAVE@aws-0-sa-east-1.pooler.supabase.com:5432/postgres"

La dirección es la misma de cargar_datos.py (Supabase → Connect → Session pooler).
Solo lee: no cambia nada en la base.
"""
import json
import os
import sys
from datetime import datetime

try:
    import psycopg2
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('Faltan librerías. Corré:  pip install psycopg2-binary openpyxl  y volvé a probar.')

# Lo cargado desde el sistema nuevo (pruebas, altas nuevas) no viene del Access: no entra en la comparación
SOLO_ACCESS = 'm.creado_por is null'

ARIAL = 'Arial'
VERDE, ROJO, GRIS = 'E2F0D9', 'FBE3E1', 'F2F2F2'
F_TIT = Font(name=ARIAL, size=14, bold=True)
F_SUB = Font(name=ARIAL, size=10, italic=True, color='555555')
F_ENC = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
F_NOR = Font(name=ARIAL, size=10)
F_NEG = Font(name=ARIAL, size=10, bold=True)
FILL_ENC = PatternFill('solid', fgColor='0B8A22')
FILL_GRIS = PatternFill('solid', fgColor=GRIS)
BORDE = Border(bottom=Side(style='thin', color='DDDDDD'))
NUM0, NUM3 = '#,##0;-#,##0;0', '#,##0.###;-#,##0.###;0'


def consultar(cur, sql, args=None):
    cur.execute(sql, args)
    return cur.fetchall()


def hoja(wb, titulo, subtitulo, encabezados, filas, anchos, formatos=None, dif=None):
    """Tabla con encabezado. dif=(col_access, col_nuevo): agrega Diferencia y Estado con fórmulas."""
    ws = wb.create_sheet(titulo)
    ws['A1'] = titulo; ws['A1'].font = F_TIT
    ws['A2'] = subtitulo; ws['A2'].font = F_SUB
    enc = list(encabezados) + (['Diferencia', 'Estado'] if dif else [])
    for j, e in enumerate(enc, 1):
        c = ws.cell(row=4, column=j, value=e)
        c.font, c.fill, c.alignment = F_ENC, FILL_ENC, Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 30
    for i, f in enumerate(filas, 5):
        for j, v in enumerate(f, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F_NOR, BORDE
            if formatos and formatos[j - 1]: c.number_format = formatos[j - 1]
        if dif:
            a, n = get_column_letter(dif[0]), get_column_letter(dif[1])
            d = ws.cell(row=i, column=len(encabezados) + 1, value=f'=ROUND({n}{i}-{a}{i},3)')
            d.font, d.border, d.number_format = F_NOR, BORDE, NUM3
            e = ws.cell(row=i, column=len(encabezados) + 2, value=f'=IF(ABS({get_column_letter(len(encabezados) + 1)}{i})<0.001,"OK","Revisar")')
            e.font, e.border, e.alignment = F_NEG, BORDE, Alignment(horizontal='center')
    for j, w in enumerate(anchos + ([13, 10] if dif else []), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    if dif and filas:
        col = get_column_letter(len(encabezados) + 2)
        from openpyxl.formatting.rule import CellIsRule
        rango = f'{col}5:{col}{4 + len(filas)}'
        ws.conditional_formatting.add(rango, CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', fgColor=VERDE)))
        ws.conditional_formatting.add(rango, CellIsRule(operator='equal', formula=['"Revisar"'], fill=PatternFill('solid', fgColor=ROJO)))
    return ws, (f"='{titulo}'!{get_column_letter(len(encabezados) + 2)}5:{get_column_letter(len(encabezados) + 2)}{max(5, 4 + len(filas))}" if dif else None)


def pedir_direccion(dada=None):
    """La dirección de la base: la variable SGRP_DB, la que se pasó (si es una de verdad) o se pide sin mostrarla."""
    import getpass
    d = (dada or '').strip()
    if d.startswith(('postgresql://', 'postgres://')) or '=' in d:
        return d
    d = os.environ.get('SGRP_DB', '').strip()
    if d: return d
    print('Pegá la dirección de la base (Supabase → Connect → Session pooler, con tu contraseña en lugar de [YOUR-PASSWORD]).')
    print('No se muestra mientras la pegás: pegala con clic derecho y tocá Enter.')
    d = getpass.getpass('Dirección: ').strip()
    if not d.startswith(('postgresql://', 'postgres://')):
        sys.exit('Eso no parece una dirección de Supabase: tiene que empezar con postgresql://')
    if '[YOUR-PASSWORD]' in d:
        sys.exit('Falta reemplazar [YOUR-PASSWORD] por la contraseña de la base.')
    return d


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) < 1: sys.exit(__doc__)
    esperado = json.load(open(args[0], encoding='utf-8'))
    dsn = pedir_direccion(args[1] if len(args) > 1 else None)
    try:
        db = psycopg2.connect(dsn, connect_timeout=20)
    except psycopg2.OperationalError as e:
        sys.exit(f'No se pudo conectar: {e}')
    db.set_session(readonly=True)
    cur = db.cursor()
    print('Leyendo la base…')

    # ------------------------------------------------------------------ conteos
    q = {
        'manifiestos': f'select count(*) from manifiestos m where {SOLO_ACCESS}',
        'manifiesto_items': f'select count(*) from manifiesto_items i join manifiestos m on m.id = i.manifiesto_id where {SOLO_ACCESS}',
        'manifiestos_nacion': 'select count(*) from manifiestos_nacion where camion_id is null and chofer_id is null',   # los del Access no tienen camión ni chofer
        'certificados': 'select count(*) from certificados',
        'generadores': 'select count(*) from generadores where id_access is not null',
        'generadores_principales': 'select count(*) from generadores_principales where id_access is not null',
        'transportistas': 'select count(*) from transportistas where id_access is not null',
        'camiones': 'select count(*) from camiones c join transportistas t on t.id = c.transportista_id where t.id_access is not null',
        'choferes': 'select count(*) from choferes c join transportistas t on t.id = c.transportista_id where t.id_access is not null and not c.creado_en_migracion',
        'operadores': 'select count(*) from operadores',
        'corrientes': 'select count(*) from corrientes where not creada_en_migracion',
    }
    nombres = {'manifiestos': 'Manifiestos', 'manifiesto_items': 'Residuos (renglones)', 'manifiestos_nacion': 'Manifiestos Nación',
               'certificados': 'Certificados de disposición final', 'generadores': 'Generadores (clientes)',
               'generadores_principales': 'Plantas (generadores principales)', 'transportistas': 'Transportistas',
               'camiones': 'Camiones', 'choferes': 'Choferes', 'operadores': 'Operadores de disposición final', 'corrientes': 'Corrientes'}
    notas = {'choferes': 'Además hay choferes que figuraban en manifiestos y no estaban en la tabla: se crearon dados de baja (ver Inconsistencias).',
             'corrientes': 'Además se crearon las corrientes usadas en manifiestos que no estaban en la tabla (ver Inconsistencias).',
             'manifiestos_nacion': 'Sin contar los que se armen desde el sistema nuevo.'}
    conteos = [(nombres[k], esperado['conteos'][k], consultar(cur, sql)[0][0], notas.get(k, '')) for k, sql in q.items()]

    # ------------------------------------------------------------------ por año / estado / unidad
    filas = consultar(cur, f"""
        select coalesce(extract(year from m.fecha)::int::text, 'sin fecha'), i.estado::text, i.um::text, count(*), coalesce(sum(i.cantidad), 0)
        from manifiesto_items i join manifiestos m on m.id = i.manifiesto_id where {SOLO_ACCESS} group by 1, 2, 3""")
    base = {f'{a}|{e}|{u}': (n, float(c)) for a, e, u, n, c in filas}
    ESTADO = {'almacenado': 'Almacenado', 'en_transito': 'En tránsito', 'disposicion_final': 'Dispuesto'}
    UNIDAD = {'KGS': 'kg', 'LTS': 'litros', 'MS3': 'm³'}
    claves = sorted(set(base) | set(esperado['residuos_por_anio']))
    por_anio = []
    for k in claves:
        a, e, u = k.split('|')
        x = esperado['residuos_por_anio'].get(k, {'residuos': 0, 'cantidad': 0})
        n, c = base.get(k, (0, 0.0))
        por_anio.append((a, ESTADO.get(e, e), UNIDAD.get(u, u), x['residuos'], n, x['cantidad'], round(c, 3)))

    man_anio = dict(consultar(cur, f"select coalesce(extract(year from m.fecha)::int::text, 'sin fecha'), count(*) from manifiestos m where {SOLO_ACCESS} group by 1"))
    man_por_anio = [(a, esperado['manifiestos_por_anio'].get(a, 0), man_anio.get(a, 0)) for a in sorted(set(man_anio) | set(esperado['manifiestos_por_anio']))]

    # ------------------------------------------------------------------ por generador
    filas = consultar(cur, f"""
        select g.id_access, g.nombre, count(i.id),
               coalesce(sum(i.cantidad) filter (where i.um = 'KGS'), 0), coalesce(sum(i.cantidad) filter (where i.um = 'LTS'), 0)
        from generadores g left join manifiestos m on m.generador_id = g.id and {SOLO_ACCESS}
        left join manifiesto_items i on i.manifiesto_id = m.id
        where g.id_access is not null group by 1, 2""")
    gen = {str(a): (nom, n, float(k), float(l)) for a, nom, n, k, l in filas}
    por_gen = []
    for ida in sorted(set(gen) | set(esperado['por_generador']), key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else 0)):
        x = esperado['por_generador'].get(ida, {'residuos': 0, 'KGS': 0, 'LTS': 0})
        nom, n, k, l = gen.get(ida, ('(no está en la base)', 0, 0.0, 0.0))
        if not x['residuos'] and not n: continue
        por_gen.append((int(ida) if ida.isdigit() else ida, nom, x['residuos'], n, x['KGS'], round(k, 3), x['LTS'], round(l, 3)))

    # ------------------------------------------------------------------ por corriente
    filas = consultar(cur, f"""
        select upper(c.codigo), count(i.id), coalesce(sum(i.cantidad) filter (where i.um = 'KGS'), 0), coalesce(sum(i.cantidad) filter (where i.um = 'LTS'), 0)
        from manifiesto_items i join manifiestos m on m.id = i.manifiesto_id join corrientes c on c.id = i.corriente_id
        where {SOLO_ACCESS} group by 1""")
    cor = {c: (n, float(k), float(l)) for c, n, k, l in filas}
    esp_cor = {}
    for c, v in esperado['por_corriente'].items():
        e = esp_cor.setdefault(c.upper(), {'residuos': 0, 'KGS': 0, 'LTS': 0})
        for kk in e: e[kk] += v.get(kk, 0)
    por_cor = []
    for c in sorted(set(cor) | set(esp_cor)):
        x = esp_cor.get(c, {'residuos': 0, 'KGS': 0, 'LTS': 0}); n, k, l = cor.get(c, (0, 0.0, 0.0))
        por_cor.append((c, x['residuos'], n, x['KGS'], round(k, 3), x['LTS'], round(l, 3)))

    # ------------------------------------------------------------------ faltantes y sobrantes
    m_base = {f'{n}|{t}' for n, t in consultar(cur, f'select m.numero, t.id_access from manifiestos m join transportistas t on t.id = m.transportista_id where {SOLO_ACCESS}')}
    m_esp = set(esperado['claves']['manifiestos'])
    tr = dict(consultar(cur, 'select id_access, nombre from transportistas where id_access is not null'))
    def legible(k):
        n, t = k.split('|'); return (f'{n[:3]}-{n[3:]}' if len(n) == 9 else n, tr.get(int(t), t))
    m_sis = {f'{n}|{t}' for n, t in consultar(cur, 'select m.numero, t.id_access from manifiestos m join transportistas t on t.id = m.transportista_id where m.creado_por is not null')}
    faltantes = [('Manifiesto', *legible(k), 'Está en el Access y no en el sistema nuevo' if k not in m_sis else
                  'Ese número se usó en el sistema nuevo (¿una prueba?): volver a aplicar las novedades y escribir REEMPLAZAR') for k in sorted(m_esp - m_base)]
    faltantes += [('Manifiesto', *legible(k), 'Está en el sistema nuevo y no en el Access') for k in sorted(m_base - m_esp)]
    n_base = {r[0] for r in consultar(cur, 'select numero from manifiestos_nacion')}
    faltantes += [('Manifiesto Nación', n, '', 'Está en el Access y no en el sistema nuevo') for n in sorted(set(esperado['claves']['manifiestos_nacion']) - n_base)]
    g_base = {r[0] for r in consultar(cur, 'select id_access from generadores where id_access is not null')}
    faltantes += [('Generador', f'Id Access {g}', '', 'Está en el Access y no en el sistema nuevo') for g in sorted(set(esperado['claves']['generadores']) - g_base)]

    # ------------------------------------------------------------------ inconsistencias y archivos
    incons_res = consultar(cur, """
        select tabla, problema, resultado, count(*) from (
          -- la carga inicial y cada migración de novedades suman casos: cada uno se cuenta una vez
          select distinct on (tabla, clave_access, problema) tabla, problema, resultado from migracion_inconsistencias
          order by tabla, clave_access, problema, corrida_id desc) x
        group by 1, 2, 3 order by 3, 4 desc""")
    RESULT = {'corregido': 'Corregido en la migración', 'migrado_sin_enlace': 'Migrado (sin el enlace)', 'revisar': 'Para revisar con Serco', 'descartado': 'Descartado'}
    incons_res = [(t, p, RESULT.get(r, r), n) for t, p, r, n in incons_res]
    incons_det = consultar(cur, """
        select tabla, clave_access, problema, resultado, detalle from (
          select distinct on (tabla, clave_access, problema) tabla, clave_access, problema, resultado, coalesce(detalle::text, '') detalle
          from migracion_inconsistencias order by tabla, clave_access, problema, corrida_id desc) x
        order by resultado, tabla, problema, clave_access""")
    incons_det = [(t, k, p, RESULT.get(r, r), d) for t, k, p, r, d in incons_det]
    archivos = consultar(cur, """
        select coalesce(carpeta, '(otra)'), count(*), count(*) filter (where encontrado), count(*) filter (where subido is not null)
        from migracion_archivos group by 1 order by 2 desc""")

    # ------------------------------------------------------------------ Excel
    wb = Workbook(); wb.remove(wb.active)
    rs = wb.create_sheet('Resumen')
    hojas = {}
    ws, hojas['Cantidades por tabla'] = hoja(wb, 'Cantidades por tabla', 'Registros en el Access y en el sistema nuevo.',
        ['Qué', 'Access', 'Sistema nuevo', 'Nota'], conteos, [34, 12, 14, 70], [None, NUM0, NUM0, None], dif=(2, 3))
    ws, hojas['Manifiestos por año'] = hoja(wb, 'Manifiestos por año', 'Cantidad de manifiestos por año de retiro.',
        ['Año', 'Access', 'Sistema nuevo'], man_por_anio, [12, 12, 14], [None, NUM0, NUM0], dif=(2, 3))
    ws, hojas['Residuos por año'] = hoja(wb, 'Residuos por año', 'Renglones de residuos y cantidad total por año, estado y unidad. La diferencia es en cantidad.',
        ['Año', 'Estado', 'Unidad', 'Renglones Access', 'Renglones sistema nuevo', 'Cantidad Access', 'Cantidad sistema nuevo'],
        por_anio, [10, 14, 9, 14, 14, 18, 18], [None, None, None, NUM0, NUM0, NUM3, NUM3], dif=(6, 7))
    ws, hojas['Por generador'] = hoja(wb, 'Por generador', 'Kilos retirados por generador (la diferencia es en kilos; los litros, al lado).',
        ['Id Access', 'Generador', 'Renglones Access', 'Renglones nuevo', 'Litros Access', 'Litros nuevo', 'Kilos Access', 'Kilos nuevo'],
        [(a, b, c, d, g, h, e, f) for a, b, c, d, e, f, g, h in por_gen], [10, 42, 12, 12, 14, 14, 16, 16],
        [None, None, NUM0, NUM0, NUM3, NUM3, NUM3, NUM3], dif=(7, 8))
    ws, hojas['Por corriente'] = hoja(wb, 'Por corriente', 'Kilos por corriente (la diferencia es en kilos; los litros, al lado).',
        ['Corriente', 'Renglones Access', 'Renglones nuevo', 'Litros Access', 'Litros nuevo', 'Kilos Access', 'Kilos nuevo'],
        [(a, b, c, f, g, d, e) for a, b, c, d, e, f, g in por_cor], [16, 12, 12, 14, 14, 16, 16],
        [None, NUM0, NUM0, NUM3, NUM3, NUM3, NUM3], dif=(6, 7))
    hoja(wb, 'Faltantes', 'Formularios que están de un lado y no del otro. Vacío = no falta nada.',
         ['Qué', 'Número', 'Transportista', 'Situación'], faltantes or [('—', 'No falta ninguno', '', '')], [20, 16, 34, 46])
    hoja(wb, 'Inconsistencias', 'Lo que la migración encontró en el Access y cómo se resolvió.',
         ['Tabla del Access', 'Problema', 'Resultado', 'Cantidad'], incons_res, [26, 70, 26, 11], [None, None, None, NUM0])
    hoja(wb, 'Inconsistencias (detalle)', 'Cada caso, con su clave en el Access.',
         ['Tabla del Access', 'Clave', 'Problema', 'Resultado', 'Detalle'], incons_det, [22, 22, 60, 24, 60])
    hoja(wb, 'Archivos', 'PDF e imágenes que el Access tenía enlazados (manifiestos firmados, certificados, habilitaciones).',
         ['Carpeta', 'Enlazados', 'Encontrados', 'Subidos al sistema'], archivos or [('—', 0, 0, 0)], [26, 13, 13, 18], [None, NUM0, NUM0, NUM0])

    # Resumen: cada control mira la columna Estado de su hoja (se recalcula si se corrige algo)
    rs['A1'] = 'Validación de la migración del Access · Gestión Serco'; rs['A1'].font = F_TIT
    rs['A2'] = (f"Access del {datetime.fromisoformat(esperado['access']['desde']):%d/%m/%Y} al {datetime.fromisoformat(esperado['access']['hasta']):%d/%m/%Y} "
                f"· base controlada el {datetime.now():%d/%m/%Y %H:%M}"); rs['A2'].font = F_SUB
    for j, e in enumerate(['Control', 'Resultado', 'Casos a revisar', 'Qué compara'], 1):
        c = rs.cell(row=4, column=j, value=e); c.font, c.fill = F_ENC, FILL_ENC
    controles = [
        ('Cantidades por tabla', 'Que no falte ningún registro: manifiestos, residuos, certificados, flota, generadores.'),
        ('Manifiestos por año', 'Que cada año tenga los mismos manifiestos.'),
        ('Residuos por año', 'Kilos, litros y metros cúbicos por año y estado (almacenado, en tránsito, dispuesto).'),
        ('Por generador', 'Lo retirado a cada generador.'),
        ('Por corriente', 'Lo retirado de cada corriente.'),
    ]
    for i, (h, que) in enumerate(controles, 5):
        rango = hojas[h][1:]
        rs.cell(row=i, column=1, value=h).font = F_NEG
        rs.cell(row=i, column=3, value=f'=COUNTIF({rango},"Revisar")').font = F_NOR
        rs.cell(row=i, column=2, value=f'=IF(C{i}=0,"OK","Revisar")').font = F_NEG
        rs.cell(row=i, column=4, value=que).font = F_NOR
    i = 5 + len(controles)
    rs.cell(row=i, column=1, value='Faltantes').font = F_NEG
    rs.cell(row=i, column=3, value=len(faltantes)).font = F_NOR
    rs.cell(row=i, column=2, value=f'=IF(C{i}=0,"OK","Revisar")').font = F_NEG
    rs.cell(row=i, column=4, value='Formularios que están en un lado y no en el otro.').font = F_NOR
    from openpyxl.formatting.rule import CellIsRule
    rs.conditional_formatting.add(f'B5:B{i}', CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', fgColor=VERDE)))
    rs.conditional_formatting.add(f'B5:B{i}', CellIsRule(operator='equal', formula=['"Revisar"'], fill=PatternFill('solid', fgColor=ROJO)))
    i += 2
    total_inc = sum(r[3] for r in incons_res)
    revisar = sum(r[3] for r in incons_res if r[2] == RESULT['revisar'])
    enl, enc, sub = (sum(r[k] for r in archivos) for k in (1, 2, 3)) if archivos else (0, 0, 0)
    for texto in [
        f'Inconsistencias del Access: {total_inc:,} casos; {revisar:,} para revisar con Serco (hoja Inconsistencias).'.replace(',', '.'),
        f'Archivos enlazados: {enl:,}; encontrados {enc:,}; subidos al sistema {sub:,} (hoja Archivos).'.replace(',', '.'),
        'Equivalencias de la migración: estado desconocido → almacenado; unidad desconocida → kilos; '
        'los marcados "ANULADO" en el chofer o la dirección se migran como anulados.',
        'No se comparan los manifiestos cargados desde el sistema nuevo (pruebas o uso real): no vienen del Access.',
    ]:
        rs.cell(row=i, column=1, value=texto).font = F_NOR; i += 1
    for col, w in zip('ABCD', [26, 12, 15, 90]): rs.column_dimensions[col].width = w

    salida = os.path.join(os.path.dirname(os.path.abspath(sys.argv[1])), f'validacion-migracion-{datetime.now():%Y-%m-%d}.xlsx')
    wb.save(salida)

    # Resumen en pantalla
    mal = [n for n, a, b, _ in conteos if a != b]
    dif_anio = [r for r in por_anio if abs(r[5] - r[6]) >= 0.001]
    dif_gen = [r for r in por_gen if abs(r[4] - r[5]) >= 0.001 or abs(r[6] - r[7]) >= 0.001]
    dif_cor = [r for r in por_cor if abs(r[3] - r[4]) >= 0.001 or abs(r[5] - r[6]) >= 0.001]
    print(f'\nCantidades por tabla: {"OK" if not mal else "diferencias en " + ", ".join(mal)}')
    print(f'Residuos por año:     {"OK" if not dif_anio else str(len(dif_anio)) + " diferencias"}')
    print(f'Por generador:        {"OK" if not dif_gen else str(len(dif_gen)) + " diferencias"}')
    print(f'Por corriente:        {"OK" if not dif_cor else str(len(dif_cor)) + " diferencias"}')
    print(f'Faltantes:            {"ninguno" if not faltantes else len(faltantes)}')
    print(f'Archivos:             {enl} enlazados, {sub} subidos')
    print(f'\nInforme: {salida}')


if __name__ == '__main__':
    main()
